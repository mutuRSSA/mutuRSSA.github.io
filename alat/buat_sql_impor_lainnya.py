"""
buat_sql_impor_lainnya.py
=========================
Membuat SATU file SQL untuk memindahkan data lama (Google Sheet) ke Supabase:
  - Risk Register + Profil Risiko RS   (Database Risk Register.xlsx)
  - Proyek FMEA                        (Database Risk Register.xlsx)
  - PDSA & Validasi Mutu               (DATABASE_MUTU_MASTER.xlsx)
  - Laporan IKP & KPC                  (Laporan IKP.xlsx, Laporan KPC.xlsx)
  - Survei Budaya Keselamatan          (Survey Budaya Keselamatan/*.xlsx)

Data mutu harian TIDAK di sini (pakai halaman impor_data_lama.html).

Cara pakai (dari folder proyek):
    python alat/buat_sql_impor_lainnya.py
Hasil: data_lama/impor_data_lainnya.sql  (berisi data pasien -> sudah di .gitignore)
Lalu jalankan isinya SEKALI di Supabase SQL Editor. Aman dijalankan ulang:
data yang sudah ada dilewati.

Skrip ini tidak berisi data; aman di-commit.
"""
import datetime, glob, hashlib, json, os, re, sys, collections, warnings
warnings.filterwarnings("ignore")
import openpyxl

AKAR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
DL = os.path.join(AKAR, "data_lama")
KELUAR = os.path.join(DL, "impor_data_lainnya.sql")
TZ = "+07:00"  # tanggal di Google Sheet = jam lokal WIB

# ---------------------------------------------------------------- pemetaan
UNIT_GANTI = {"Unit Rawat Inap": "Unit Rawat Inap - Truntum"}   # keputusan Komite Mutu
UNIT_TERKAIT_IKP = {                                              # singkatan lama -> nama unit
    "ugd": "Unit Gawat Darurat", "igd": "Unit Gawat Darurat", "rawat jalan": "Unit Rawat Jalan",
    "rawat inap": "Unit Rawat Inap - Truntum", "ipsrs": "Unit Pemeliharaan Sarana Rumah Sakit",
    "rekam medis": "Unit Rekam Medis", "farmasi": "Unit Farmasi", "laboratorium": "Unit Laboratorium",
    "radiologi": "Unit Radiologi", "gizi": "Unit Gizi", "icu": "Unit Rawat Intensif", "ok": "Unit Kamar Operasi",
    "vk": "Unit Kamar Bersalin", "perinatologi": "Unit Perinatologi", "kasir": "Unit Kasir",
}
# Judul indikator PDSA lama yang tidak persis sama dengan Profil Indikator
PDSA_JUDUL_KE_FORM = {
    "waktu penyediaan alat steril": "Waktu_Sterilisasi_Alat",
    "pemberi pelayanan kegawatdaruratan yang bersertifikasi atls/btls/acls/ppgd": "Anggota_UGD_Terlatih",
}
SKALA_DAMPAK = {"tidak signifikan": 1, "minor": 2, "moderat": 3, "moderate": 3, "mayor": 4, "major": 4, "katastropik": 5}
SKALA_PROB = {"sangat jarang": 1, "jarang": 2, "mungkin": 3, "sering": 4, "sangat sering": 5}


def grading(d, p):  # matriks yang sama dengan ikp.html
    if d == 5: return "Merah" if p >= 2 else "Kuning"
    if d == 4: return "Merah" if p >= 3 else "Kuning"
    if d == 3: return "Merah" if p >= 4 else ("Kuning" if p == 3 else "Hijau")
    if d == 2: return "Kuning" if p == 5 else ("Hijau" if p >= 3 else "Biru")
    return "Hijau" if p >= 4 else "Biru"


# ---------------------------------------------------------------- utilitas
norm = lambda s: re.sub(r"\s+", " ", str(s or "")).strip().lower()
teks = lambda v: "" if v is None else (v.strip() if isinstance(v, str) else str(v))


def angka(v):
    if v in (None, ""): return None
    try:
        f = float(str(v).replace(",", "."))
        return int(f) if f.is_integer() else f
    except ValueError:
        return None


def tgl_tz(v):   # timestamptz
    if isinstance(v, datetime.datetime): return v.strftime("%Y-%m-%dT%H:%M:%S") + TZ
    if isinstance(v, datetime.date): return v.strftime("%Y-%m-%dT00:00:00") + TZ
    return None


def tgl_lokal(v):  # timestamp tanpa zona (waktu_insiden)
    if isinstance(v, datetime.datetime): return v.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(v, datetime.date): return v.strftime("%Y-%m-%dT00:00:00")
    return None


def tgl_ddmmyyyy(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime("%d/%m/%Y")
    return teks(v)


def uuid_dari(*bagian):  # UUID tetap (idempoten) dari kunci lama
    h = hashlib.md5("|".join(str(b) for b in bagian).encode()).hexdigest()
    return f"{h[:8]}-{h[8:12]}-{h[12:16]}-{h[16:20]}-{h[20:]}"


def baca(file, sheet):
    ws = openpyxl.load_workbook(file, read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    h = [teks(x) for x in rows[0]]
    return [dict(zip(h, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]


catatan = []
unit = lambda u: UNIT_GANTI.get(teks(u), teks(u))

# ---------------------------------------------------------------- MASTER
master = os.path.join(DL, "DATABASE_MUTU_MASTER.xlsx")
ind = baca(master, "Setup_Indikator")
judul_ke_form = {norm(r["Judul Indikator"]): teks(r["ID_Form"]) for r in ind if r.get("ID_Form")}
judul_ke_form.update(PDSA_JUDUL_KE_FORM)


def cari_form(judul):
    n = norm(judul)
    if n in judul_ke_form: return judul_ke_form[n]
    import difflib
    c = difflib.get_close_matches(n, list(judul_ke_form), n=1, cutoff=0.6)
    if c: return judul_ke_form[c[0]]
    catatan.append(f"PDSA/Validasi: indikator '{teks(judul)[:60]}' tidak dikenali (id_indikator dikosongkan)")
    return None


# ---------------------------------------------------------------- RISIKO
rr = os.path.join(DL, "Database Risk Register.xlsx")
risiko_terbaru = {}
for r in baca(rr, "DB_RISK_REGISTER"):
    rid = teks(r["ID Risiko"])
    if rid and (rid not in risiko_terbaru or (r["Timestamp"] or datetime.datetime.min) >= (risiko_terbaru[rid]["Timestamp"] or datetime.datetime.min)):
        risiko_terbaru[rid] = r
profil = {teks(p["ID Risiko"]): p for p in baca(rr, "DB_PROFIL_RS")}
RISIKO = []
for rid, r in risiko_terbaru.items():
    d, f, k = angka(r["Dampak (D)"]), angka(r["Frekuensi (F)"]), angka(r["Kontrol (K)"])
    p = profil.get(rid, {})
    RISIKO.append({
        "id_risiko": rid, "tahun": r["Timestamp"].year if isinstance(r["Timestamp"], datetime.datetime) else 2026,
        "unit_kerja": unit(r["Unit Kerja"]), "risiko_teridentifikasi": teks(r["Risiko Teridentifikasi"]),
        "penyebab": teks(r["Penyebab"]), "dampak": d, "frekuensi": f, "dampak_timbul": teks(r["Dampak Timbul"]),
        "penanganan_saat_ini": teks(r["Penanganan Saat Ini"]), "kontrol": k,
        "skor_risiko": angka(r["Skor Risiko (DxF)"]), "skor_akhir": angka(r["Skor Akhir"]), "status": "terkirim",
        "is_profil_rs": bool(p), "kategori_risiko": teks(p.get("Kategori Risiko")) or None,
        "rencana_penanganan": teks(p.get("Rencana Penanganan")) or None, "periode_program": teks(p.get("Periode Program")) or None,
        "pic_penanganan": teks(p.get("PIC Penanganan")) or None, "realisasi": teks(p.get("Realisasi")) or None,
        "rencana_lanjutan": teks(p.get("Rencana Pencegahan Lanjutan")) or None, "hasil_pemantauan": teks(p.get("Hasil Pemantauan")) or None,
        "status_penanganan": teks(p.get("Status Penanganan")) or None,
    })
for pid in profil:
    if pid not in risiko_terbaru: catatan.append(f"Profil Risiko RS: ID {pid} tidak ada di register (dilewati)")

# ---------------------------------------------------------------- FMEA
proyek = baca(rr, "Setup_FMEA_Proyek")
detail = collections.defaultdict(list)
for r in baca(rr, "DB_FMEA_Daftar"):
    detail[teks(r["ID Proyek"])].append([
        teks(r["Langkah Proses"]), teks(r["Mode Kegagalan"]), teks(r["Akibat"]), angka(r["S"]), teks(r["Penyebab"]),
        angka(r["O"]), teks(r["Kontrol"]), angka(r["D"]), angka(r["RPN"]), teks(r["Rencana Tindakan"]), teks(r["PIC"]),
        tgl_ddmmyyyy(r["Target Tanggal"]), "", ""])   # 2 kolom baru: Tindak Lanjut / Realisasi, Status
FMEA = []
for p in proyek:
    pid = teks(p["ID Proyek"])
    FMEA.append({"id_proyek": uuid_dari("fmea", pid), "tanggal_dibuat": tgl_tz(p["Tanggal Dibuat"]),
                 "nama_proses": teks(p["Nama Proses"]), "ketua_tim": teks(p["Ketua Tim"]),
                 "sumber_risiko": teks(p["Sumber Risiko"]), "status_proyek": "Aktif", "tabel_data": detail.get(pid, [])})
for pid in detail:
    if pid not in {teks(p["ID Proyek"]) for p in proyek}:
        catatan.append(f"FMEA: {len(detail[pid])} baris detail milik proyek {pid} yang tidak ada di daftar proyek (dilewati)")

# ---------------------------------------------------------------- PDSA
PDSA = []
for r in baca(master, "DB_PDSA"):
    fase = "SELESAI" if norm(r["Status"]) == "selesai" else "PLAN"
    if teks(r.get("Act_Keputusan")): fase = "SELESAI"
    elif teks(r.get("Study_Evaluasi")): fase = "STUDY"
    elif teks(r.get("Do_Pelaksanaan")) and fase != "SELESAI": fase = "DO"
    analisa, tl, kepala = teks(r["Analisa_Akar_Masalah"]), teks(r["Tindak_Lanjut_Plan"]), teks(r["Kepala_Unit"])
    PDSA.append({
        "id_pdsa": uuid_dari("pdsa", r["ID_PDSA"]), "tanggal_dibuat": tgl_tz(r["Tanggal"]), "unit_kerja": unit(r["Unit"]),
        "judul_indikator": re.sub(r"\s+", " ", teks(r["Indikator"])).strip(), "id_indikator": cari_form(r["Indikator"]),
        "periode_analisis": teks(r["Periode"]), "fase_saat_ini": fase,
        "plan_data": {"tim": kepala, "masalah": analisa, "akar": analisa, "tujuan": "", "action": [{"act": tl, "pic": kepala, "tgl": ""}]},
        "do_data": {"pelaksanaan": teks(r.get("Do_Pelaksanaan")), "hambatan": ""},
        "study_data": {"data_sebelum": "", "data_sesudah": "", "analisis": teks(r.get("Study_Evaluasi"))},
        "act_data": {"keputusan": teks(r.get("Act_Keputusan")), "tindak_lanjut": ""},
    })

# ---------------------------------------------------------------- VALIDASI
VALIDASI = []
for r in baca(master, "DB_VALIDASI"):
    VALIDASI.append({
        "id_validasi": uuid_dari("validasi", r["ID Validasi"]), "tanggal_validasi": tgl_tz(r["Tgl Validasi"]),
        "unit_kerja": unit(r["Unit"]), "id_indikator": cari_form(r["Indikator"]), "judul_indikator": teks(r["Indikator"]),
        "bulan": teks(r["Bulan"]), "tahun": angka(r["Tahun"]), "populasi_n_besar": angka(r["Populasi (N)"]),
        "sampel_n_kecil": angka(r["Sampel (n)"]), "capaian_a": angka(r["Capaian Total (A)"]),
        "capaian_b": angka(r["Capaian Sampel (B)"]), "akurasi": angka(r["Akurasi (%)"]),
        "status_validasi": teks(r["Status Validasi"]), "nama_validator": teks(r["Validator"]),
    })

# ---------------------------------------------------------------- IKP
INSIDEN = []
KOLOM_INVESTIGASI_IKP = {
    "Tipe & Sub Tipe Insiden": "tipe_insiden", "Tipe Harm": "tipe_harm", "Masalah CMP": "masalah_cmp", "Masalah SDP": "masalah_sdp",
    "Penyebab Langsung": "penyebab_langsung", "Akar Masalah": "akar_masalah", "Root Cause (Pilihan)": "root_cause_select",
    "Orang Terlibat": "orang_terlibat", "Proses/Fase Pelayanan": "proses_pelayanan", "Faktor Kontributor": "faktor_kontributor",
    "Mitigasi Pasien": "mitigasi_pasien", "Mitigasi Fasyankes": "mitigasi_fasyankes", "Mitigasi Faktor Terkait": "mitigasi_faktor",
    "Cara Mendeteksi": "cara_mendeteksi", "Dampak Fasyankes": "dampak_fasyankes", "Rekomendasi": "rekomendasi",
    "Tindakan Akan Dilakukan": "tindakan_dilakukan", "PJ Tindakan": "pj_tindakan", "Perbaikan Kepada Pasien": "perbaikan_pasien",
    "Perbaikan Kepada Fasyankes": "perbaikan_fasyankes", "Risiko Kepada Pasien": "risiko_pasien", "Risiko Kepada Petugas": "risiko_petugas",
    "Risiko Untuk Lingkungan": "risiko_lingkungan", "Risiko Faktor Mempengaruhi": "risiko_faktor", "Pembelajaran": "pembelajaran",
    "Safety Alert": "safety_alert"}
for r in baca(os.path.join(DL, "Laporan IKP.xlsx"), "Data"):
    d = SKALA_DAMPAK.get(norm(r["Dampak"]), angka(r["Dampak"]) or 1)
    p = SKALA_PROB.get(norm(r["Probabilitas"]), angka(r["Probabilitas"]) or 1)
    terkait = teks(r["Unit Terkait"])
    unit_ikp = UNIT_TERKAIT_IKP.get(norm(terkait), terkait or "Tidak Diketahui")
    inv = {v: teks(r.get(k)) for k, v in KOLOM_INVESTIGASI_IKP.items() if teks(r.get(k))}
    INSIDEN.append({
        "id_insiden": uuid_dari("ikp", tgl_tz(r["Timestamp"]), teks(r["Insiden"])),
        "waktu_lapor": tgl_tz(r["Timestamp"]), "waktu_insiden": tgl_lokal(r["Tgl & Waktu insiden"]),
        "unit_pelapor": unit_ikp, "lokasi_kejadian": unit_ikp, "nama_pasien": "-", "no_rm": "-",
        "jenis_insiden": teks(r["Jenis Insiden"]),
        "kronologi": f"Insiden: {teks(r['Insiden'])}\n\nKronologis:\n{teks(r['Kronologis'])}",
        "tindakan_segera": teks(r["Tindak Lanjut"]), "pelaksana_tindakan": teks(r["Dilakukan Oleh"]),
        "skor_dampak": d, "skor_probabilitas": p, "grading_risiko": grading(d, p),
        "nama_pelapor": "Tidak tercatat (data lama)",
        "detail_spesifik": {
            "umur_hari": teks(r["Umur (Hari)"]), "umur_bulan": teks(r["Umur (Bulan)"]), "umur_tahun": teks(r["Umur (Tahun)"]),
            "jenis_kelamin": teks(r["Jenis Kelamin"]), "penanggung_biaya": teks(r["Penanggung Biaya"]),
            "tgl_pelayanan": tgl_ddmmyyyy(r["Tanggal pelayanan"]), "terjadi_pada_pasien": teks(r["Terjadi Pada"]),
            "unit_penyebab": terkait, "pelapor_pertama": teks(r["Pelapor Pertama"]),
            "pernah_terjadi_sebelumnya": teks(r["Pernah terjadi?"]), "tindakan_pencegahan": teks(r["Tindakan Pencegahan"]),
            "menyangkut_pasien": teks(r["Menyangkut Pasien"]), "tempat_insiden": teks(r["Tempat Insiden"]),
            "grading_lama": teks(r["Grading Risiko"]), "sumber": "Google Sheet (data lama)"},
        "investigasi_komite": ({**inv, "tanggal_investigasi": None} if inv else None),
        "status_investigasi": "Selesai" if inv else None,
    })

# ---------------------------------------------------------------- KPC
for r in baca(os.path.join(DL, "Laporan KPC.xlsx"), "Data_KPC"):
    inv = {k2: teks(r.get(k1)) for k1, k2 in {
        "Penyebab Langsung": "penyebab_langsung", "Akar Masalah": "akar_masalah", "Rekomendasi": "rekomendasi",
        "PJ Rekomendasi": "pj_rekomendasi", "Tgl rekomendasi": "tgl_rekomendasi", "Tindakan Akan Dilakukan": "tindakan_dilakukan",
        "PJ Tindakan": "pj_tindakan", "Tgl Tindakan": "tgl_tindakan"}.items() if teks(r.get(k1))}
    if "tgl_rekomendasi" in inv: inv["tgl_rekomendasi"] = tgl_ddmmyyyy(r.get("Tgl rekomendasi"))
    if "tgl_tindakan" in inv: inv["tgl_tindakan"] = tgl_ddmmyyyy(r.get("Tgl Tindakan"))
    status_lama = norm(r.get("Status KPC"))
    INSIDEN.append({
        "id_insiden": uuid_dari("kpc", tgl_tz(r["Timestamp"]), teks(r["Deskripsi KPC"])),
        "waktu_lapor": tgl_tz(r["Timestamp"]), "waktu_insiden": tgl_lokal(r["Tgl & Waktu Penemuan"]),
        "unit_pelapor": unit(r["Unit Pelapor"]), "lokasi_kejadian": teks(r["Lokasi / Unit"]),
        "nama_pasien": "-", "no_rm": "-", "jenis_insiden": "KPC", "kronologi": teks(r["Deskripsi KPC"]),
        "tindakan_segera": teks(r["Tindakan yang Dilakukan"]), "pelaksana_tindakan": "Pelapor (Penemu KPC)",
        "skor_dampak": 0, "skor_probabilitas": 0, "grading_risiko": "Biru", "nama_pelapor": teks(r["Nama Pelapor"]) or "Anonim",
        "detail_spesifik": {"tindakan_lanjutan": teks(r.get("Tindakan Akan Dilakukan")), "pj_tindakan": teks(r.get("PJ Tindakan")),
                            "target_tgl_selesai": tgl_ddmmyyyy(r.get("Tgl Tindakan")), "foto_bukti": teks(r.get("Link Bukti Foto")),
                            "status_lama": teks(r.get("Status KPC")), "sumber": "Google Sheet (data lama)"},
        "investigasi_komite": ({**inv, "tanggal_investigasi": None} if inv else None),
        "status_investigasi": "Selesai" if status_lama in ("closed", "selesai", "tutup") else None,
    })

# ---------------------------------------------------------------- SURVEI
SURVEI = []
for f in glob.glob(os.path.join(DL, "Survey Budaya Keselamatan", "*.xlsx")):
    for r in baca(f, openpyxl.load_workbook(f, read_only=True).sheetnames[0]):
        jawaban = {k: (v if isinstance(v, (int, float)) else teks(v)) for k, v in r.items() if k not in ("Timestamp", "unit_kerja", "profesi") and k}
        SURVEI.append({"id_survey": uuid_dari("survei", tgl_tz(r["Timestamp"]), teks(r["unit_kerja"]), teks(r["profesi"])),
                       "timestamp": tgl_tz(r["Timestamp"]), "unit_kerja": unit(r["unit_kerja"]), "profesi": teks(r["profesi"]),
                       "jawaban_survey": jawaban})

# ---------------------------------------------------------------- SQL
def js(obj):
    s = json.dumps(obj, ensure_ascii=False, default=str)
    if "$DATA$" in s: sys.exit("Data mengandung penanda $DATA$ - ganti penanda di skrip.")
    return f"$DATA${s}$DATA$::jsonb"


unit_dicek = sorted({x["unit_kerja"] for x in RISIKO + PDSA + VALIDASI + SURVEI} |
                    {x["unit_pelapor"] for x in INSIDEN if x["jenis_insiden"] != "KPC"})

sql = f"""-- =====================================================================
-- IMPOR DATA LAMA (Google Sheet) -> Supabase
-- Dibuat oleh alat/buat_sql_impor_lainnya.py pada {datetime.datetime.now():%d/%m/%Y %H:%M}
-- BERISI DATA PASIEN: jangan dibagikan / di-commit.
--
-- Jalankan SEKALI di Supabase SQL Editor. Satu transaksi: jika ada error,
-- tidak ada yang berubah. Aman dijalankan ulang (data yang sudah ada dilewati).
--
-- Isi: {len(RISIKO)} risiko ({sum(1 for r in RISIKO if r['is_profil_rs'])} profil RS), {len(FMEA)} proyek FMEA,
--      {len(PDSA)} PDSA, {len(VALIDASI)} validasi, {len(INSIDEN)} insiden (IKP+KPC), {len(SURVEI)} survei.
-- =====================================================================

begin;

-- 0. Pastikan semua nama unit ada di daftar unit aplikasi
do $$
declare v_hilang text;
begin
  select string_agg(u, ', ') into v_hilang
    from unnest(array{json.dumps(unit_dicek, ensure_ascii=False).replace('"', "'")}::text[]) u
   where not exists (select 1 from public.master_unit m where m.nama_unit = u);
  if v_hilang is not null then
    raise exception 'Unit berikut belum ada di Pengaturan Unit aplikasi: %. Tambahkan dulu (atau sesuaikan nama), lalu jalankan ulang.', v_hilang;
  end if;
end $$;

create temp table _hasil (tabel text, dimasukkan int) on commit drop;

-- 1. RISK REGISTER + PROFIL RISIKO RS
with x as (select * from jsonb_to_recordset({js(RISIKO)}) as t(
    id_risiko text, tahun int, unit_kerja text, risiko_teridentifikasi text, penyebab text, dampak int, frekuensi int,
    dampak_timbul text, penanganan_saat_ini text, kontrol int, skor_risiko int, skor_akhir int, status text,
    is_profil_rs boolean, kategori_risiko text, rencana_penanganan text, periode_program text, pic_penanganan text,
    realisasi text, rencana_lanjutan text, hasil_pemantauan text, status_penanganan text)),
ins as (insert into public.data_risiko (id_risiko, tahun, unit_kerja, risiko_teridentifikasi, penyebab, dampak, frekuensi,
    dampak_timbul, penanganan_saat_ini, kontrol, skor_risiko, skor_akhir, status, is_profil_rs, kategori_risiko,
    rencana_penanganan, periode_program, pic_penanganan, realisasi, rencana_lanjutan, hasil_pemantauan, status_penanganan)
  select * from x on conflict (id_risiko) do nothing returning 1)
insert into _hasil select 'data_risiko', count(*) from ins;

-- 2. FMEA
with x as (select * from jsonb_to_recordset({js(FMEA)}) as t(
    id_proyek uuid, tanggal_dibuat timestamptz, nama_proses text, ketua_tim text, sumber_risiko text, status_proyek text, tabel_data jsonb)),
ins as (insert into public.data_fmea (id_proyek, tanggal_dibuat, nama_proses, ketua_tim, sumber_risiko, status_proyek, tabel_data)
  select * from x where not exists (select 1 from public.data_fmea f where f.id_proyek = x.id_proyek
                                       or (f.nama_proses = x.nama_proses and f.tanggal_dibuat = x.tanggal_dibuat))
  returning 1)
insert into _hasil select 'data_fmea', count(*) from ins;

-- 3. PDSA
with x as (select * from jsonb_to_recordset({js(PDSA)}) as t(
    id_pdsa uuid, tanggal_dibuat timestamptz, unit_kerja text, judul_indikator text, id_indikator text,
    periode_analisis text, fase_saat_ini text, plan_data jsonb, do_data jsonb, study_data jsonb, act_data jsonb)),
ins as (insert into public.data_pdsa (id_pdsa, tanggal_dibuat, unit_kerja, judul_indikator, id_indikator, periode_analisis,
    fase_saat_ini, plan_data, do_data, study_data, act_data)
  select * from x where not exists (select 1 from public.data_pdsa d where d.id_pdsa = x.id_pdsa
       or (d.unit_kerja = x.unit_kerja and d.periode_analisis = x.periode_analisis and d.judul_indikator = x.judul_indikator))
  returning 1)
insert into _hasil select 'data_pdsa', count(*) from ins;

-- 4. VALIDASI MUTU
with x as (select * from jsonb_to_recordset({js(VALIDASI)}) as t(
    id_validasi uuid, tanggal_validasi timestamptz, unit_kerja text, id_indikator text, judul_indikator text, bulan text,
    tahun int, populasi_n_besar int, sampel_n_kecil int, capaian_a numeric, capaian_b numeric, akurasi numeric,
    status_validasi text, nama_validator text)),
ins as (insert into public.data_validasi (id_validasi, tanggal_validasi, unit_kerja, id_indikator, judul_indikator, bulan, tahun,
    populasi_n_besar, sampel_n_kecil, capaian_a, capaian_b, akurasi, status_validasi, nama_validator)
  select * from x where not exists (select 1 from public.data_validasi v where v.id_validasi = x.id_validasi
       or (v.unit_kerja = x.unit_kerja and v.judul_indikator = x.judul_indikator and v.bulan = x.bulan and v.tahun = x.tahun))
  returning 1)
insert into _hasil select 'data_validasi', count(*) from ins;

-- 5. INSIDEN (IKP + KPC)
with x as (select * from jsonb_to_recordset({js(INSIDEN)}) as t(
    id_insiden uuid, waktu_lapor timestamptz, waktu_insiden timestamp, unit_pelapor text, lokasi_kejadian text,
    nama_pasien text, no_rm text, jenis_insiden text, kronologi text, tindakan_segera text, pelaksana_tindakan text,
    skor_dampak int, skor_probabilitas int, grading_risiko text, nama_pelapor text, detail_spesifik jsonb,
    investigasi_komite jsonb, status_investigasi text)),
ins as (insert into public.data_insiden (id_insiden, waktu_lapor, waktu_insiden, unit_pelapor, lokasi_kejadian, nama_pasien,
    no_rm, jenis_insiden, kronologi, tindakan_segera, pelaksana_tindakan, skor_dampak, skor_probabilitas, grading_risiko,
    nama_pelapor, detail_spesifik, investigasi_komite, status_investigasi)
  select id_insiden, waktu_lapor, waktu_insiden, unit_pelapor, lokasi_kejadian, nama_pasien, no_rm, jenis_insiden, kronologi,
         tindakan_segera, pelaksana_tindakan, skor_dampak, skor_probabilitas, grading_risiko, nama_pelapor, detail_spesifik,
         investigasi_komite, coalesce(status_investigasi, 'Menunggu Review')
    from x where not exists (select 1 from public.data_insiden i where i.id_insiden = x.id_insiden
       or (i.waktu_insiden = x.waktu_insiden and i.jenis_insiden = x.jenis_insiden and i.kronologi = x.kronologi))
  returning 1)
insert into _hasil select 'data_insiden', count(*) from ins;

-- 6. SURVEI BUDAYA
with x as (select * from jsonb_to_recordset({js(SURVEI)}) as t(
    id_survey uuid, "timestamp" timestamptz, unit_kerja text, profesi text, jawaban_survey jsonb)),
ins as (insert into public.data_survey_budaya (id_survey, "timestamp", unit_kerja, profesi, jawaban_survey)
  select * from x where not exists (select 1 from public.data_survey_budaya s where s.id_survey = x.id_survey)
  returning 1)
insert into _hasil select 'data_survey_budaya', count(*) from ins;

-- Ringkasan (lihat hasil di bawah)
select tabel, dimasukkan from _hasil order by tabel;

commit;
"""
open(KELUAR, "w", encoding="utf-8").write(sql)
print(f"OK -> {os.path.relpath(KELUAR, AKAR)}")
print(f"   risiko {len(RISIKO)} (profil RS {sum(1 for r in RISIKO if r['is_profil_rs'])}), FMEA {len(FMEA)} proyek "
      f"({sum(len(f['tabel_data']) for f in FMEA)} baris), PDSA {len(PDSA)}, validasi {len(VALIDASI)}, "
      f"insiden {len(INSIDEN)} (IKP {sum(1 for i in INSIDEN if i['jenis_insiden'] != 'KPC')}, KPC {sum(1 for i in INSIDEN if i['jenis_insiden'] == 'KPC')}), survei {len(SURVEI)}")
print("   unit yang dipakai:", ", ".join(unit_dicek))
for c in catatan: print("   CATATAN:", c)
